import itertools
import unittest
from datetime import datetime
from io import BytesIO
from openpyxl import Workbook
from duties import assign_duties
from finance import read_members, make_report
from openpyxl import load_workbook


def members(n=10):
    return [{'姓名': str(i), '角色': '队员', '报名时间': datetime(2026, 1, 1),
             '校区': '学院路' if i < 3 else '沙河', '是否新人': i % 2 == 0,
             '愿意负责': i % 3 == 0, '职能意向': ['摄影', '装备', '队记', '后勤', '队医']}
            for i in range(n)]


class AssignmentTests(unittest.TestCase):
    def test_constraints_and_reproducibility(self):
        source = members(27)
        source += [{'姓名': role, '角色': role} for role in ('领队', '副领队', '财务')]
        result = assign_duties(source)
        self.assertEqual(result, assign_duties(source))
        self.assertNotIn('职能', source[0])
        for duty in ('后勤', '队医', '装备', '摄影', '队记'):
            group = [m for m in result if m['职能'] == duty]
            self.assertLessEqual(len(group), 6)
            leaders = [m for m in group if m['职能负责']]
            self.assertEqual(len(leaders), 1)
            if duty in ('后勤', '队医', '装备'):
                self.assertEqual(leaders[0]['校区'], '学院路')
        self.assertEqual(result[-1]['职能'], '财务')
        self.assertEqual(result[-2]['职能'], '')

    def test_global_objectives_against_enumeration(self):
        source = members(5)
        duties = ['后勤', '队医', '装备', '摄影', '队记']
        for i, m in enumerate(source):
            m['校区'] = '学院路'
            m['职能意向'] = duties[i:] + duties[:i]
        source.append({**source[0], '姓名': 'extra', '愿意负责': True, '是否新人': False})
        def score(mapping, leaders):
            return (sum(not source[i]['愿意负责'] for i in leaders),
                    sum(source[i]['是否新人'] for i in leaders),
                    sum(source[i]['职能意向'].index(d) for i, d in enumerate(mapping)))
        best = None
        for mapping in itertools.product(duties, repeat=6):
            groups = [[i for i, d in enumerate(mapping) if d == duty] for duty in duties]
            if any(not g or len(g) > 2 for g in groups):
                continue
            for leaders in itertools.product(*groups):
                value = score(mapping, leaders)
                best = value if best is None else min(best, value)
        result = assign_duties(source)
        self.assertEqual(score([m['职能'] for m in result],
                               [i for i, m in enumerate(result) if m['职能负责']]), best)

    def test_invalid_inputs(self):
        with self.assertRaisesRegex(ValueError, '至少需要 5'):
            assign_duties(members(4))
        source = members()
        source[0]['校区'] = '沙河'
        with self.assertRaisesRegex(ValueError, '至少需要 3'):
            assign_duties(source)
        source[0]['是否新人'] = None
        with self.assertRaisesRegex(ValueError, '缺少校区或是否新人'):
            assign_duties(source)

    def test_equipment_discount(self):
        for pole, lamp in itertools.product((False, True), repeat=2):
            member = {'姓名': '测试', '角色': '财务', '登山杖': pole, '头灯': lamp}
            wb = load_workbook(BytesIO(make_report([member], 1000, 2, 12)))
            params, report = wb['费用参数'], wb['财务报表']
            self.assertEqual([params[f'B{i}'].value for i in (9, 10, 11)], [12, 8, 0.8])
            self.assertEqual(report['F2'].value, f"='费用参数'!B9*'费用参数'!B11*{int(pole)}")
            self.assertEqual(report['G2'].value, f"='费用参数'!B10*'费用参数'!B11*{int(lamp)}")
            self.assertEqual(round(params['B9'].value * params['B11'].value * pole, 2), 9.6 if pole else 0)
            self.assertEqual(round(params['B10'].value * params['B11'].value * lamp, 2), 6.4 if lamp else 0)

    def test_import_fields(self):
        wb = Workbook()
        ws = wb.active
        ws.append(['姓名', '报名时间', '所在校区', '是否新人'])
        ws.append(['甲', datetime(2026, 1, 1), '学院路校区', '○否'])
        stream = BytesIO()
        wb.save(stream)
        stream.seek(0)
        result = read_members(stream)
        self.assertEqual(result[0]['校区'], '学院路')
        self.assertIs(result[0]['是否新人'], False)


if __name__ == '__main__':
    unittest.main()
