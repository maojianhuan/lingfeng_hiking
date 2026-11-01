import re
import unittest
from io import BytesIO
from unittest.mock import Mock

from openpyxl import load_workbook

from app import OutingApp
from finance import ATTENDING, WITHDRAW_DAY_BEFORE, WITHDRAW_SAME_DAY, make_report


def member(name, state=ATTENDING, role='队员'):
    return {'姓名': name, '角色': role, '出队状态': state, '登山杖': True, '头灯': True}


def numeric_formula(workbook, sheet, address):
    """Evaluate only arithmetic and parameter references in these exported fee cells."""
    value = workbook[sheet][address].value
    if not isinstance(value, str) or not value.startswith('='):
        return value
    expression = re.sub(r"'([^']+)'!([A-Z]+\d+)",
                        lambda m: str(numeric_formula(workbook, m[1], m[2])), value[1:])
    return eval(expression, {'__builtins__': {}}, {})


class SettlementTests(unittest.TestCase):
    def test_withdrawal_fees_and_bus_conservation(self):
        people = [member('正常'), member('提前一天', WITHDRAW_DAY_BEFORE),
                  member('当天', WITHDRAW_SAME_DAY), member('领队', role='领队')]
        wb = load_workbook(BytesIO(make_report(people, 200, 2, 12, [('公共药品', 40), ('公共食品', 20)])))
        ws = wb['财务报表']
        self.assertEqual([numeric_formula(wb, ws.title, f'H{r}') for r in range(2, 6)], [60, 20, 60, 60])
        self.assertEqual(sum(numeric_formula(wb, ws.title, f'H{r}') for r in range(2, 6)), 200)
        for r in (3, 4):
            for col in ('D', 'F', 'G'):
                self.assertEqual(numeric_formula(wb, ws.title, f'{col}{r}'), 0)
        for r in range(2, 6):
            self.assertEqual(numeric_formula(wb, ws.title, f'E{r}'), 2)
            self.assertEqual(numeric_formula(wb, ws.title, f'I{r}'), 6)
            self.assertEqual(numeric_formula(wb, ws.title, f'K{r}'), 10)
            self.assertEqual(numeric_formula(wb, ws.title, f'L{r}'), 5)
            self.assertEqual(ws[f'M{r}'].value, 0)
            self.assertEqual(ws[f'N{r}'].value, f'=ROUND(C{r}-SUM(D{r}:M{r}),2)')
        # Advance withdrawal: 100 - insurance 2 - bus 20 - radio 6 - shared supplies 15 = 57.
        self.assertEqual(100 - sum(numeric_formula(wb, ws.title, f'{col}3')
                                   for col in ('D', 'E', 'F', 'G', 'H', 'I', 'K', 'L', 'M')), 57)
        # Day-of withdrawal: same fees with a 60-yuan bus share, refund 17.
        self.assertEqual(100 - sum(numeric_formula(wb, ws.title, f'{col}4')
                                   for col in ('D', 'E', 'F', 'G', 'H', 'I', 'K', 'L', 'M')), 17)
        ws['M4'] = 10
        self.assertEqual(100 - sum(numeric_formula(wb, ws.title, f'{col}4')
                                   for col in ('D', 'E', 'F', 'G', 'H', 'I', 'K', 'L', 'M')), 7)
        self.assertTrue(wb.calculation.fullCalcOnLoad)
        self.assertEqual(wb['公共支出']['A2'].value, '公共药品')
        self.assertEqual(wb['公共支出']['B2'].value, 40)
        self.assertNotIn('当天', str(list(wb['职能安排'].values)))

    def test_financing_without_extra_expenses(self):
        wb = load_workbook(BytesIO(make_report([member('正常')], 0, 0, 12)))
        ws = wb['财务报表']
        self.assertEqual(ws['D2'].value, '=IF(OR(B2="领队",B2="副领队"),0,\'费用参数\'!B7)')
        self.assertEqual(ws['K1'].value, '融资')
        self.assertEqual(ws['K2'].value, 0)
        self.assertEqual(ws['L2'].value, '=ROUND(C2-SUM(D2:K2),2)')
        self.assertEqual(ws.auto_filter.ref, 'A1:L2')
        self.assertEqual(ws['L3'].value, '=SUM(L2:L2)')

    def test_multiple_advance_withdrawals(self):
        people = [member('甲', WITHDRAW_DAY_BEFORE), member('乙', WITHDRAW_DAY_BEFORE), member('丙')]
        wb = load_workbook(BytesIO(make_report(people, 100, 0, 12)))
        self.assertEqual([numeric_formula(wb, '财务报表', f'H{r}') for r in (2, 3, 4)], [20, 20, 60])

    def test_invalid_settlements_fail(self):
        with self.assertRaisesRegex(ValueError, '至少需要一名'):
            make_report([member('退出', WITHDRAW_DAY_BEFORE)], 100, 0, 12)
        with self.assertRaisesRegex(ValueError, '固定车费总额'):
            make_report([member('退出', WITHDRAW_DAY_BEFORE), member('出队')], 10, 0, 12)
        with self.assertRaisesRegex(ValueError, '未知出队状态'):
            make_report([member('甲', '错误状态')], 100, 0, 12)
        with self.assertRaisesRegex(ValueError, '公共支出'):
            make_report([member('甲')], 100, 0, 12, [('药品', -1)])

    def test_selection_and_export_keep_withdrawals_only_in_settlement(self):
        app = OutingApp.__new__(OutingApp)
        app.members = [member('同名'), member('同名'), member('当天'), member('提前两天')]
        app.member_list = Mock()
        app.member_list.curselection.return_value = (0, 1, 2)
        app.withdrawals = {1: WITHDRAW_DAY_BEFORE, 2: WITHDRAW_SAME_DAY}
        self.assertEqual(app.selected_members(), [app.members[0]])
        app.get_assignment = Mock(return_value=[{**app.members[0], '角色': '领队'}])
        app.cost_values = Mock(return_value=(100, 0))
        app.expenses = Mock()
        app.expenses.get_children.return_value = ()
        app.save_report = Mock()
        app.export_finance()
        content = app.save_report.call_args.args[0]
        rows = list(load_workbook(BytesIO(content))['财务报表'].values)
        self.assertEqual([row[9] for row in rows[1:4]], [ATTENDING, WITHDRAW_DAY_BEFORE, WITHDRAW_SAME_DAY])
        self.assertEqual([row[1] for row in rows[1:4]], ['领队', '退出', '退出'])
        self.assertNotIn('提前两天', [row[0] for row in rows[1:4]])


if __name__ == '__main__':
    unittest.main()
