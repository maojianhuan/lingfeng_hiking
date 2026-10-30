from datetime import datetime
import re
from duties import DUTIES, duty_rows
from io import BytesIO
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def read_members(source):
    wb = load_workbook(source, read_only=True, data_only=True)
    try:
        for sheet in wb:
            rows = list(sheet.values)
            for pos, row in enumerate(rows):
                headers = [str(v).strip() if v is not None else '' for v in row]
                headers = [re.sub(r'^\d+\.\s*', '', h).strip() for h in headers]
                aliases = {'你的姓名': '姓名', '提交答卷时间': '报名时间',
                           '职能志愿': '职能意向', '所在校区': '校区',
                           '年级（不区分大一，研一等称呼，选择入学年级）': '入学年份',
                           '年级': '入学年份',
                           '身份证号（仅用于购买保险）': '身份证号',
                           '联系方式（电话）': '手机号',
                           '是否愿意担任职能负责（建议新人积极担任）': '是否愿意担任职能负责'}
                headers = [aliases.get(h, h) for h in headers]
                if '姓名' not in headers or '报名时间' not in headers:
                    continue
                ni, ti = headers.index('姓名'), headers.index('报名时间')
                members = []
                for number, values in enumerate(rows[pos + 1:], pos + 2):
                    if all(v is None for v in values):
                        continue
                    name, time = values[ni], values[ti]
                    if not name or not time:
                        raise ValueError(f'第 {number} 行缺少姓名或报名时间')
                    if not isinstance(time, datetime):
                        try:
                            text = str(time).strip()
                            time = datetime.strptime(text, '%Y/%m/%d %H:%M:%S') if '/' in text else datetime.fromisoformat(text)
                        except ValueError:
                            raise ValueError(f'第 {number} 行报名时间须为 Excel 日期或 YYYY-MM-DD HH:MM:SS')
                    preferences = []
                    if '职能意向' in headers:
                        raw = values[headers.index('职能意向')]
                        clean = re.sub(r'（[^）]*）|\([^)]*\)', '', str(raw or ''))
                        preferences = [v.strip() for v in re.split(r'[、,，>＞→;；\n]', clean) if v.strip()]
                        if len(set(preferences)) != len(preferences) or any(v not in DUTIES for v in preferences):
                            raise ValueError(f'第 {number} 行职能意向须为不重复的职能名称')
                    willing = False
                    if '是否愿意担任职能负责' in headers:
                        raw = values[headers.index('是否愿意担任职能负责')]
                        raw = str(raw).strip().lstrip('○') if raw is not None else None
                        if raw not in (None, '', '是', '否'):
                            raise ValueError(f'第 {number} 行是否愿意担任职能负责须填写是或否')
                        willing = raw == '是'
                    year = None
                    if '入学年份' in headers:
                        raw_year = values[headers.index('入学年份')]
                        text_year = str(raw_year).strip()
                        if not re.fullmatch(r'\d{4}', text_year):
                            raise ValueError(f'第 {number} 行入学年份须为四位年份')
                        year = int(text_year)
                    campus = None
                    if '校区' in headers:
                        raw_campus = values[headers.index('校区')]
                        campus = str(raw_campus).strip() if raw_campus is not None else None
                        if campus == '学院路校区':
                            campus = '学院路'
                    newcomer = None
                    if '是否新人' in headers:
                        raw_newcomer = values[headers.index('是否新人')]
                        text_newcomer = str(raw_newcomer).strip().lstrip('○') if raw_newcomer is not None else ''
                        if text_newcomer not in ('', '是', '否'):
                            raise ValueError(f'第 {number} 行是否新人须填写是或否')
                        if text_newcomer:
                            newcomer = text_newcomer == '是'
                    members.append({'姓名': str(name).strip(), '报名时间': time, '入学年份': year,
                                    '职能意向': preferences, '愿意负责': willing,
                                    '校区': campus, '是否新人': newcomer,
                                    '身份证号': values[headers.index('身份证号')] if '身份证号' in headers else None,
                                    '手机号': values[headers.index('手机号')] if '手机号' in headers else None,
                                    '登山杖': '需要租借的装备' in headers and '登山杖' in str(values[headers.index('需要租借的装备')]),
                                    '头灯': '需要租借的装备' in headers and '头灯' in str(values[headers.index('需要租借的装备')])})
                return sorted(members, key=lambda m: m['报名时间'])
        raise ValueError('找不到“姓名”和“报名时间”列。问卷汇总表不包含个人记录，请导出答卷明细或手动填写。')
    finally:
        wb.close()


def make_report(members, bus, radios, radio_price):
    if not members:
        raise ValueError('请至少选择一名成员')
    if bus < 0 or radios < 0 or radio_price < 0 or int(radios) != radios:
        raise ValueError('费用须非负，对讲机数量须为非负整数')
    wb = Workbook()
    ws = wb.active
    ws.title = '财务报表'
    params = wb.create_sheet('费用参数')
    for row in [('项目', '数值'), ('包车费用', bus), ('对讲机数量', radios),
                ('对讲机单价', radio_price), ('出队人数', len(members)),
                ('默认预收', 100), ('领队费', 10), ('保险', 2), ('登山杖原价', 12), ('头灯原价', 8), ('装备租赁折扣（暂全员适用）', 0.8)]:
        params.append(row)
    ws.append(['姓名', '角色', '预收金额', '领队费', '保险', '登山杖费', '头灯费', '包车分摊', '对讲机分摊', '剩余金额'])
    for r, m in enumerate(members, 2):
        ws.append([m['姓名'], m['角色'], "='费用参数'!B6",
                   f'=IF(OR(B{r}="领队",B{r}="副领队"),0,\'费用参数\'!B7)',
                   "='费用参数'!B8",
                   f"='费用参数'!B9*'费用参数'!B11*{int(m['登山杖'])}",
                   f"='费用参数'!B10*'费用参数'!B11*{int(m['头灯'])}",
                   "='费用参数'!B2/'费用参数'!B5", "='费用参数'!B3*'费用参数'!B4/'费用参数'!B5",
                   f'=ROUND(C{r}-SUM(D{r}:I{r}),2)'])
    duties = wb.create_sheet('职能安排')
    for row in duty_rows(members):
        duties.append(row)
    for r in range(2, duties.max_row + 1):
        duties.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
    end = len(members) + 1
    ws.append(['合计', '', *[f'=SUM({get_column_letter(c)}2:{get_column_letter(c)}{end})' for c in range(3, 11)]])
    ws.cell(end + 3, 1, '剩余为正数应退，为负数应补；分摊保留计算精度，个人余额四舍五入到分，可能有分币尾差。')
    for sheet in wb:
        sheet.freeze_panes = 'A2'
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='246B55')
        for col in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 18
    duties.freeze_panes = 'A2'
    for column in ('A', 'C', 'E'):
        duties.column_dimensions[column].width = 14
    for column in ('B', 'D', 'F'):
        duties.column_dimensions[column].width = 32
    for row in duties:
        duties.row_dimensions[row[0].row].height = 36
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='center')
    for r in range(2, duties.max_row + 1):
        duties.cell(r, 1).font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, max_row=end + 1):
        for cell in row:
            if 3 <= cell.column <= 10:
                cell.number_format = '0.00;[Red]-0.00'
    ws.auto_filter.ref = f'A1:J{end}'
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def make_insurance_report(members):
    if not members:
        raise ValueError('请至少选择一名成员')
    wb = Workbook()
    ws = wb.active
    ws.title = '保险名单'
    ws.append(['姓名', '身份证号', '手机号'])
    for r, member in enumerate(members, 2):
        for col, field in enumerate(['姓名', '身份证号', '手机号'], 1):
            value = member.get(field)
            if value is None or not str(value).strip():
                raise ValueError(f'第 {r - 1} 位出队成员缺少{field}，请补充后导出')
            if field == '身份证号' and isinstance(value, (int, float)):
                raise ValueError(f'第 {r - 1} 位出队成员身份证号为数值，可能已丢失精度，请从原始记录以文本重新填写')
            if isinstance(value, float) and value.is_integer():
                value = int(value)
            cell = ws.cell(r, col)
            cell.value = str(value).strip()
            cell.data_type = 's'
            cell.number_format = '@'
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='246B55')
    for column, width in [('A', 18), ('B', 26), ('C', 20)]:
        ws.column_dimensions[column].width = width
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
